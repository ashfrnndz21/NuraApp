"""Create one GitHub Issue per backlog story from docs/feature-backlog.xlsx.

Usage: GH_REPO=org/nura python scripts/backlog_to_issues.py [--tier T1] [--dry-run]
Requires the gh CLI authenticated. Idempotent: skips stories whose ID already appears in an open or closed issue title.
"""
import argparse, os, subprocess, sys
from openpyxl import load_workbook

HIGH = {"E00", "E12", "E04", "E16", "E19", "E05", "E06", "E20"}
MEDIUM = {"E02", "E03", "E09", "E11", "E17", "E18", "E21"}

def risk(epic: str, story: str) -> str:
    s = story.lower()
    if any(k in s for k in ("key", "consent", "audit", "medic", "dose", "red flag", "send", "outbound", "migration", "escalat")):
        return "risk:high"
    if epic in HIGH:
        return "risk:high"
    return "risk:medium" if epic in MEDIUM else "risk:low"

def existing_titles() -> set[str]:
    out = subprocess.run(["gh", "issue", "list", "--state", "all", "--limit", "1000", "--json", "title", "-q", ".[].title"], capture_output=True, text=True, check=True).stdout
    return set(out.split("\n"))

def main() -> None:
    ap = argparse.ArgumentParser(); ap.add_argument("--tier", default=None); ap.add_argument("--dry-run", action="store_true"); a = ap.parse_args()
    ws = load_workbook("docs/feature-backlog.xlsx", data_only=True)["Backlog"]
    have = set() if a.dry_run else existing_titles()
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for (sid, epic, epic_name, story, persona, tier, pri, size, deps, acc, *_rest) in rows:
        if a.tier and tier != a.tier: continue
        title = f"{sid} {story[:80]}"
        if any(t.startswith(sid) for t in have): continue
        body = f"""**Epic** {epic} {epic_name}
**Persona** {persona} · **Tier** {tier} · **Priority** {pri} · **Size** {size}

**Story**
{story}

**Acceptance**
{acc}

**Depends on**
{deps or 'none'}

**Spec** see docs/00-MASTER-BUILD-SPEC.md section 0 for the document that covers this epic.
"""
        labels = ["story", tier, epic, risk(epic, story)]
        cmd = ["gh", "issue", "create", "--title", title, "--body", body] + sum([["--label", l] for l in labels], [])
        print(" ".join(cmd[:5]), labels)
        if not a.dry_run:
            subprocess.run(cmd, check=True)

if __name__ == "__main__":
    if not os.environ.get("GH_REPO"):
        print("Set GH_REPO=org/repo (or run inside the repo with gh configured)", file=sys.stderr)
    main()
